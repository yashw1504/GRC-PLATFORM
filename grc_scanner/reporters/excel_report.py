"""Excel workbook export with summary, compliance, and severity worksheets."""

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelReport:
    HEADERS = [
        "check_id", "name", "status", "severity", "risk_score", "priority",
        "cve_id", "cvss_score", "category", "scanner_name", "target_type",
        "target_asset", "description", "evidence", "mapped_controls",
        "remediation", "recommendation",
    ]
    SEVERITIES = ("Critical", "High", "Medium", "Low", "Info")

    @staticmethod
    def _value(value):
        if value is None:
            return ""
        if isinstance(value, (dict, list, tuple)):
            value = json.dumps(value, ensure_ascii=False, default=str)
        # Excel cells have a 32,767-character limit.
        return str(value)[:32767]

    @classmethod
    def generate(cls, findings, compliance_scores, overall_risk_score=0,
                 output_file="output/report.xlsx"):
        Path(output_file).parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()

        summary = workbook.active
        summary.title = "Summary"
        summary.append(["Metric", "Value"])
        summary.append(["Overall Risk Score", overall_risk_score])
        summary.append(["Total Findings", len(findings)])
        for severity in cls.SEVERITIES:
            summary.append([
                f"{severity} Findings",
                sum(1 for finding in findings if finding.severity.lower() == severity.lower()),
            ])

        compliance = workbook.create_sheet("Compliance")
        compliance.append(["Framework", "Score (%)"])
        for framework, score in sorted(compliance_scores.items()):
            compliance.append([framework, score if score is not None else "Not assessed"])

        all_findings = workbook.create_sheet("Findings")
        all_findings.append(cls.HEADERS)
        for finding in findings:
            record = finding.to_dict()
            all_findings.append([cls._value(record.get(header)) for header in cls.HEADERS])

        for severity in cls.SEVERITIES:
            sheet = workbook.create_sheet(severity)
            sheet.append(cls.HEADERS)
            for finding in findings:
                if finding.severity.lower() == severity.lower():
                    record = finding.to_dict()
                    sheet.append([cls._value(record.get(header)) for header in cls.HEADERS])

        for sheet in workbook.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E79")
            for column_cells in sheet.columns:
                width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 48)
                sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

        workbook.save(output_file)
        return output_file

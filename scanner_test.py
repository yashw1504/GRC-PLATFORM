"""Isolated smoke tests for scanner behavior against a local mock HTTP server."""

from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from threading import Thread
from pathlib import Path
from tempfile import TemporaryDirectory
import importlib.util
import subprocess
import unittest
from unittest.mock import patch

from grc_scanner.scanners.web.web_scanner import WebScanner
from grc_scanner.scanners.api.api_scanner import APIScanner
from grc_scanner.engine.deduplicator import FindingDeduplicator
from grc_scanner.engine.compliance_engine import ComplianceEngine
from grc_scanner.engine.finding import Finding
from grc_scanner.engine.finding_normalizer import FindingNormalizer
from grc_scanner.engine.orchestrator import ScannerOrchestrator
from grc_scanner.engine.risk_engine import RiskEngine
from grc_scanner.checks import coverage_summary
from grc_scanner.reporters.html_report import HtmlReport
from grc_scanner.integrations.nmap_wrapper import NmapWrapper
from grc_scanner.scanners.vulnerability.vulnerability_scanner import VulnerabilityScanner
from grc_scanner.scanners.sbom.sbom_scanner import SBOMScanner

if importlib.util.find_spec("openpyxl"):
    from grc_scanner.reporters.excel_report import ExcelReport
else:
    ExcelReport = None


class VulnerableMockHandler(BaseHTTPRequestHandler):
    def do_GET(self):
        body = b"<html><body><h1>Mock vulnerable site</h1></body></html>"
        self.send_response(200)
        self.send_header("Content-Type", "text/html")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def log_message(self, *_args):
        pass


class WebScannerMockServerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.server = ThreadingHTTPServer(("127.0.0.1", 0), VulnerableMockHandler)
        cls.thread = Thread(target=cls.server.serve_forever, daemon=True)
        cls.thread.start()
        cls.target = f"http://127.0.0.1:{cls.server.server_port}/"

    @classmethod
    def tearDownClass(cls):
        cls.server.shutdown()
        cls.server.server_close()
        cls.thread.join(timeout=2)

    def test_detects_http_and_missing_security_headers(self):
        scanner = WebScanner()
        # The mock only serves HTTP. Skip the separate TLS socket probe so the
        # test remains local and deterministic.
        scanner._check_tls_version = lambda _target: []

        findings = scanner.scan(self.target)
        failed_checks = {f.check_id for f in findings if f.status == "fail"}

        self.assertIn("https_enforced", failed_checks)
        self.assertIn("hsts_present", failed_checks)
        self.assertIn("csp_present", failed_checks)
        self.assertIn("x_frame_options", failed_checks)
        self.assertIn("x_content_type_options", failed_checks)

    def test_api_scanner_reports_observed_response_properties_only(self):
        findings = APIScanner().scan(self.target)
        check_ids = {finding.check_id for finding in findings}

        self.assertIn("api_https_enforced", check_ids)
        self.assertIn("api_content_type", check_ids)
        self.assertNotIn("api_no_authentication", check_ids)


class CorePipelineTests(unittest.TestCase):
    def test_normalizer_uses_cvss_and_emits_priority(self):
        finding = Finding(
            check_id="vulnerability_found",
            name="Test CVE",
            status="FAIL",
            severity="high",
            raw_data={"cve": "CVE-2026-12345", "cvss_score": 9.4},
        )

        normalized = FindingNormalizer.normalize([finding])[0]

        self.assertEqual(normalized.status, "fail")
        self.assertEqual(normalized.cve_id, "CVE-2026-12345")
        self.assertEqual(normalized.risk_score, 9)
        self.assertEqual(normalized.priority, "P1 - Immediate")

    def test_deduplicator_merges_cve_matches_and_scanner_names(self):
        first = Finding(
            check_id="vulnerability_found", name="CVE finding", status="fail",
            severity="High", risk_score=8, cve_id="CVE-2026-12345", target_asset="app:1",
            scanner_name="nuclei", evidence="package app",
        )
        second = Finding(
            check_id="dependency_vulnerability", name="Same CVE", status="fail",
            severity="Critical", risk_score=10, cve_id="cve-2026-12345", target_asset="app:1",
            scanner_name="grype", evidence="package app",
        )

        unique = FindingDeduplicator.deduplicate([first, second])

        self.assertEqual(len(unique), 1)
        self.assertEqual(unique[0].severity, "Critical")
        self.assertIn("grype", unique[0].scanner_name)
        self.assertIn("nuclei", unique[0].scanner_name)

    def test_orchestrator_converts_scanner_exception_to_finding(self):
        class BrokenScanner:
            name = "broken"

            def scan(self, _target):
                raise RuntimeError("tool unavailable")

        orchestrator = ScannerOrchestrator()
        orchestrator.logger.exception = lambda *_args, **_kwargs: None
        findings = orchestrator.run([BrokenScanner()], "local")

        self.assertEqual(len(findings), 1)
        self.assertEqual(findings[0].status, "error")
        self.assertEqual(findings[0].category, "Scanner Error")

    def test_tool_timeout_returns_empty_list(self):
        timeout = subprocess.TimeoutExpired(cmd="nmap", timeout=1)
        with patch.object(NmapWrapper, "is_available", return_value=True), patch(
            "grc_scanner.integrations.nmap_wrapper.subprocess.run",
            side_effect=timeout,
        ):
            self.assertEqual(NmapWrapper.scan("127.0.0.1"), [])

    def test_scanner_error_does_not_dilute_overall_risk_score(self):
        finding = Finding(
            check_id="risk", name="High risk", status="fail", severity="High",
            risk_score=8,
        )
        error = Finding(
            check_id="scanner_error", name="Unavailable", status="error",
            severity="Info", risk_score=0,
        )

        self.assertEqual(RiskEngine.overall_score([finding, error]), 80)

    def test_nuclei_cve_and_cvss_are_normalized(self):
        item = {
            "template-id": "CVE-2026-12345",
            "matched-at": "https://service.example/v1/items",
            "info": {
                "name": "Example vulnerability",
                "severity": "high",
                "classification": {"cvss-score": 8.7},
            },
        }
        with patch("grc_scanner.scanners.vulnerability.vulnerability_scanner.NucleiWrapper.is_available", return_value=True), patch(
            "grc_scanner.scanners.vulnerability.vulnerability_scanner.NucleiWrapper.scan",
            return_value=[item],
        ):
            findings = VulnerabilityScanner().scan("https://service.example")

        self.assertEqual(findings[0].cve_id, "CVE-2026-12345")
        self.assertEqual(findings[0].cvss_score, 8.7)
        self.assertEqual(findings[0].risk_score, 9)

    def test_grype_finding_extracts_cve_and_cvss(self):
        finding = SBOMScanner()._create_finding(
            {
                "id": "CVE-2026-12345",
                "severity": "Critical",
                "cvss": [{"metrics": {"baseScore": 9.8}}],
            },
            {"name": "sample", "version": "1.0"},
        )

        self.assertEqual(finding.cve_id, "CVE-2026-12345")
        self.assertEqual(finding.cvss_score, 9.8)

    def test_unmapped_frameworks_are_not_reported_as_full_compliance(self):
        scores = ComplianceEngine().calculate_scores([])
        self.assertTrue(scores)
        self.assertTrue(all(score is None for score in scores.values()))

    def test_check_coverage_catalog_surfaces_gaps_and_count_mismatch(self):
        summary = coverage_summary()

        self.assertEqual(summary["prompt_expected_total"], 625)
        self.assertEqual(summary["prompt_category_count_sum"], 625)
        self.assertEqual(summary["detailed_file_count_sum"], 605)
        self.assertGreater(summary["unmapped_check_ids"], 0)
        self.assertGreater(summary["mapped_ids_missing_owasp"], 0)
        self.assertFalse(summary["definitions_complete"])

    def test_html_report_escapes_untrusted_finding_content(self):
        finding = Finding(
            check_id="test", name="<script>alert(1)</script>", status="fail",
            severity="High", evidence="<img src=x onerror=alert(1)>",
        )
        with TemporaryDirectory() as directory:
            output = Path(directory) / "report.html"
            HtmlReport.generate([finding], {}, 10, str(output))
            content = output.read_text(encoding="utf-8")

        self.assertNotIn("<script>", content)
        self.assertIn("&lt;script&gt;", content)

    @unittest.skipIf(ExcelReport is None, "openpyxl is installed in CI through requirements")
    def test_excel_report_has_eight_worksheets(self):
        finding = Finding(
            check_id="test", name="Sample", status="fail", severity="High",
        )
        with TemporaryDirectory() as directory:
            output = Path(directory) / "report.xlsx"
            ExcelReport.generate([finding], {"PCI": 90}, 20, str(output))
            from openpyxl import load_workbook
            workbook = load_workbook(output, read_only=True)
            sheetnames = workbook.sheetnames
            workbook.close()

        self.assertEqual(len(sheetnames), 8)
        self.assertEqual(sheetnames[0], "Summary")


if __name__ == "__main__":
    unittest.main()
